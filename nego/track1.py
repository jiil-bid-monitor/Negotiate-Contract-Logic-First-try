"""트랙 1 — 1억 미만 공고를 사내 NAS 폴더로 (2026-10-28 사용자 결정: 1억 미만 / 이상 두 갈래).

공고 하나에 폴더 하나를 만들고 담당자가 검토만 하면 되게 자료를 채운다:

    <root>/입찰/2026/1028_R26BK01790000_공고명/
      00_원문/                 첨부 원본 + 텍스트/(읽은 글)
      01_공고요약.pdf (+ .html) 기본 정보·일정·참가자격·확인 필요·과업 요약·제출 방법·주요 조항·투찰 참고·첨부 목록
      02_제출서류_체크리스트.xlsx 담당자가 상태를 적는 표 — 이미 있으면 덮어쓰지 않음
      일정.ics                 마감 일정(하루 전 알림)
      추출결과.json            자동으로 뽑은 값 전부(원문 위치 포함)
    <root>/입찰/공고관리대장.xlsx  공고 한 줄씩 — 시스템 칸만 고치고 담당자 칸(참여 여부·결과·메모)은 그대로

판정·일정·금액은 지금 수집기의 규칙 판정을 그대로 쓰고, 과업 요약·제출서류·제출 방법·주요 조항만 Claude API가 원문에서 뽑는다
(사용자 결정: 규칙 + Claude API). 키가 없으면 규칙 부분만 만든다. 평가 배점·점수는 다루지 않는다(2026-10-14 사용자 결정).
투찰 금액 결정과 전자서명은 담당자 몫 — 폴더는 참고 자료만.

    python -m nego.track1 --root "\\\\NAS\\공유\\입찰자료" --days 1
"""

from __future__ import annotations

import argparse
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Callable

from . import attachments, bidcalc, bidrecommend, qualify
from .config import AppConfig, ConfigError, load_config, redact
from .http_client import ApiError, DataGoKrClient

log = logging.getLogger(__name__)

DEFAULT_MAX_BUDGET = 100_000_000  # 1억 미만 = 트랙 1 (금액은 수집기 기준: 추정가격, 없으면 배정예산 그대로)
LEDGER_NAME = "공고관리대장.xlsx"
SYSTEM_COLUMNS = ["공고번호", "차수", "등록일", "공고명", "발주기관", "업무", "유형", "금액", "금액 종류", "마감", "자격", "지역",
                  "공동수급", "확인 필요", "폴더"]
STAFF_COLUMNS = ["참여 여부", "담당", "결과", "메모"]


# ── 트랙 판정·폴더 ──────────────────────────────────────────────

def is_track1(c, max_budget: float = DEFAULT_MAX_BUDGET) -> bool:
    """후보로 남은 본공고 중 금액이 1억 미만인 것. 금액을 모르는 공고는 트랙 1로 보내지 않는다(기존 시스템에서 검토)."""
    n = c.notice
    return bool(getattr(c, "is_candidate", True)) and n.kind == "본공고" and n.budget is not None and n.budget < max_budget


_UNSAFE = re.compile(r'[\\/:*?"<>|\r\n\t]')


def folder_for(root: Path, c, now: datetime | None = None) -> Path:
    """<root>/입찰/<게시 연도>/<게시 월일>_<공고번호>_<공고명 앞 40자>. 윈도우 경로 길이를 생각해 공고명을 자른다."""
    posted = getattr(c.schedule, "posted_at", None) or now or datetime.now()
    title = re.sub(r"\s+", " ", _UNSAFE.sub("_", c.notice.title)).strip(" ._")[:40].rstrip(" ._")
    return root / "입찰" / f"{posted:%Y}" / f"{posted:%m%d}_{c.notice.notice_no}_{title}"


# ── Claude API 추출 ─────────────────────────────────────────────

SYSTEM_PROMPT = """당신은 공공 입찰 공고 문서(입찰공고문·제안요청서·과업지시서·규격서)에서 담당자가 검토할 사실만 뽑는 도우미입니다.
규칙:
- 원문에 적힌 내용만 씁니다. 추측하거나 일반 상식으로 채우지 않습니다. 원문에 없으면 그 항목을 빈 목록으로 둡니다.
- 모든 항목의 src에 근거 위치를 "파일명 · 항목 번호/제목"처럼 적습니다(파일명은 "=== 파일명 ===" 줄에 있습니다).
- 평가 배점·점수·정량평가 계산은 다루지 않습니다(평가 항목 표를 옮기지 마세요).
- 담당자 이름·개인 전화번호·개인 이메일은 쓰지 않습니다.
- 한국어로, 짧고 쉬운 문장으로 씁니다."""

USER_PROMPT = """아래 공고의 첨부 원문에서 다음을 뽑아 JSON으로 답하세요.

1. task_summary: 과업 요약 4~8줄 — 목적, 대상 시설·장소·규모, 해야 할 일(설계·제작·설치 등), 납품물·기간. 한 줄에 하나.
2. submissions: 제출해야 하는 것과 그 일시·방법 — 예: 입찰참가신청서류(이메일, 6/27 09:00~11:00), 제안서(방문 제출), 전자입찰(나라장터).
   when은 원문 표기 그대로, how는 전자입찰/이메일/방문/우편/기타 중 원문에 맞는 것, where는 제출처(기관·부서·주소. 개인 연락처 제외).
3. documents: 제출서류 목록 — 원문의 번호·순서대로. name(서류명), form(서식 번호, 없으면 ""), qty(부수, 없으면 ""),
   stage(언제 내는지: 참가신청/입찰/제안서/낙찰 후 등), prep(준비 방식: 공통서류=사업자등록증·법인등기·나라장터 등록증처럼 회사가 늘 갖춘 서류,
   발급=관공서·기관에서 새로 떼야 하는 서류, 서식 작성=공고 첨부 서식에 적어 내는 서류, 직접 작성=제안서·산출내역처럼 새로 만드는 문서,
   해당 시=공동수급·대리인 등 조건부), note(원본대조필·날인·유효기간 같은 주의, 없으면 "").
4. clauses: 놓치면 손해 보는 조항 — 실격·무효 사유, 감점, 가격(예정가격 대비 하한 등), 계약 기한, 비용 부담·하자 책임, 권리 이전,
   보안, 하도급, 대금·지급 조건. kind는 그중 하나, text는 한두 문장.
5. proposal_outline: 제안서 작성 목차가 원문에 있으면 장(chapter)과 하위 항목(items) — 없으면 빈 목록.

공고명: {title}
발주기관: {institution}
유형: {category} / {award}

## 첨부 원문{cut}
{text}"""

_ITEM = {"type": "object", "additionalProperties": False}
OUTPUT_SCHEMA = {
    "type": "object", "additionalProperties": False,
    "required": ["task_summary", "submissions", "documents", "clauses", "proposal_outline"],
    "properties": {
        "task_summary": {"type": "array", "items": {**_ITEM, "required": ["text", "src"],
                                                    "properties": {"text": {"type": "string"}, "src": {"type": "string"}}}},
        "submissions": {"type": "array", "items": {**_ITEM, "required": ["what", "when", "how", "where", "src"], "properties": {
            "what": {"type": "string"}, "when": {"type": "string"}, "how": {"type": "string"},
            "where": {"type": "string"}, "src": {"type": "string"}}}},
        "documents": {"type": "array", "items": {**_ITEM, "required": ["name", "form", "qty", "stage", "prep", "note", "src"],
                                                 "properties": {
            "name": {"type": "string"}, "form": {"type": "string"}, "qty": {"type": "string"}, "stage": {"type": "string"},
            "prep": {"type": "string", "enum": ["공통서류", "발급", "서식 작성", "직접 작성", "해당 시"]},
            "note": {"type": "string"}, "src": {"type": "string"}}}},
        "clauses": {"type": "array", "items": {**_ITEM, "required": ["kind", "text", "src"], "properties": {
            "kind": {"type": "string", "enum": ["실격·무효", "감점", "가격", "계약", "비용·책임", "권리", "보안", "하도급",
                                                "대금·지급", "기타"]},
            "text": {"type": "string"}, "src": {"type": "string"}}}},
        "proposal_outline": {"type": "array", "items": {**_ITEM, "required": ["chapter", "items"], "properties": {
            "chapter": {"type": "string"}, "items": {"type": "array", "items": {"type": "string"}}}}},
    },
}


@dataclass
class LlmSettings:
    model: str
    effort: str = "medium"
    max_text_chars: int = 60_000  # 제출서류 절은 공고문 뒤쪽에 있는 일이 많아 유사도 판정(2만 자)보다 넉넉히
    max_tokens: int = 16_000

    @classmethod
    def from_env(cls) -> "LlmSettings | None":
        from .llm_similarity import LlmConfig

        base = LlmConfig.from_env()
        if base is None:
            return None
        return cls(model=base.model, effort=base.effort,
                   max_text_chars=int(os.environ.get("TRACK1_MAX_TEXT_CHARS", "60000")))


def empty_extraction(note: str) -> dict:
    return {"task_summary": [], "submissions": [], "documents": [], "clauses": [], "proposal_outline": [],
            "source": "규칙", "note": note}


def extract_with_llm(client: Any, settings: LlmSettings, c) -> dict:
    """공고 1건 → OUTPUT_SCHEMA 모양 dict(+ source·model·error). 실패해도 예외를 올리지 않는다."""
    import anthropic

    from .llm_similarity import FALLBACK_BETA

    text = (c.attachment_text or "").strip()
    if not text:
        return empty_extraction("첨부 원문을 읽지 못해 자동 추출을 건너뜀")
    cut = len(text) > settings.max_text_chars
    prompt = USER_PROMPT.format(
        title=c.notice.title, institution=c.notice.demand_institution or c.notice.notice_institution or "",
        category=c.category or "", award=c.notice.award_method or "",
        cut=f" (앞 {settings.max_text_chars:,}자만)" if cut else "", text=text[: settings.max_text_chars])
    try:
        response = client.beta.messages.create(
            model=settings.model, max_tokens=settings.max_tokens, system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": prompt}],
            output_config={"effort": settings.effort, "format": {"type": "json_schema", "schema": OUTPUT_SCHEMA}},
            betas=[FALLBACK_BETA], fallbacks="default",
        )
    except anthropic.APIStatusError as err:
        return empty_extraction(f"Claude API 오류({err.status_code}) — 규칙 판정만 넣음")
    except anthropic.APIConnectionError as err:
        return empty_extraction(f"Claude API 연결 실패({err}) — 규칙 판정만 넣음")
    if response.stop_reason in ("refusal", "max_tokens"):
        return empty_extraction(f"Claude 응답이 끝나지 않음({response.stop_reason}) — 규칙 판정만 넣음")
    body = next((b.text for b in response.content if getattr(b, "type", None) == "text"), "")
    try:
        data = json.loads(body)
    except json.JSONDecodeError:
        return empty_extraction("Claude 응답을 읽지 못함 — 규칙 판정만 넣음")
    data.update(source="Claude", model=getattr(response, "model", None) or settings.model,
                note=f"원문 앞 {settings.max_text_chars:,}자만 읽음 — 뒤쪽 내용은 원문 확인" if cut else "")
    return data


def rule_extraction(c, note: str) -> dict:
    """Claude 없이 — 과업 요약은 첨부의 '사업 범위' 목록(웹앱 과업 유사 팝업과 같은 규칙)."""
    from .business_scope import extract_scope_items

    out = empty_extraction(note)
    out["task_summary"] = [{"text": item, "src": "첨부 '사업 범위' 목록"}
                           for item in extract_scope_items(c.task_text or c.attachment_text)[:10]]
    out["excerpts"] = {"documents": excerpt_sections(c.attachment_text, "documents", 1),
                       "notes": excerpt_sections(c.attachment_text, "notes", 3)}
    return out


# 키가 없을 때 — 제출서류·유의사항 절을 원문 그대로 발췌(담당자가 원문을 뒤지지 않게, 10/28 필수 자료 요청)
_MARK = r"(?:\d{1,2}\s*[.)]|\(\d{1,2}\)|[가-하]\s*[.)]|[IVXⅠ-Ⅹ]+\s*\.|[□■○●◆◇▶◈※]|제\s*\d+\s*조)?\s*"
EXCERPT_HEADINGS = {
    "documents": re.compile(r"^\s*" + _MARK + r"(?:입찰\s*)?(?:제출\s*서류|구비\s*서류|입찰\s*참가\s*(?:신청|등록)\s*서류)"),
    "notes": re.compile(r"^\s*" + _MARK + r"(?:기타\s*(?:유의\s*)?사항|유의\s*사항|입찰\s*무효|하도급|제한\s*사항)"),
}
_NEXT_HEADING = re.compile(r"^\s*(?:\d{1,2}\s*\.|[IVXⅠ-Ⅹ]+\s*\.|제\s*\d+\s*조|[□■◆◇])\s*\S")
_FILE_LINE = re.compile(r"^=== (.+?) ===$")
MAX_EXCERPT_LINES = 45


def excerpt_sections(text: str, kind: str, limit: int = 1) -> list[dict]:
    """원문에서 그 절(제목 줄 + 다음 큰 제목 전까지, 최대 45줄)을 최대 `limit`개. 공고문 파일을 먼저 본다."""
    head_re = EXCERPT_HEADINGS[kind]
    files: list[tuple[str, list[str]]] = [("", [])]
    for line in (text or "").splitlines():
        m = _FILE_LINE.match(line.strip())
        if m:
            files.append((m.group(1), []))
        else:
            files[-1][1].append(line)
    files.sort(key=lambda f: "공고" not in f[0])
    found: list[dict] = []
    for name, lines in files:
        i = 0
        while i < len(lines) and len(found) < limit:
            line = lines[i]
            if len(line.strip()) > 40 or not head_re.match(line):
                i += 1
                continue
            body, j = [line.strip()], i + 1
            while j < len(lines) and len(body) < MAX_EXCERPT_LINES:
                nxt = lines[j]
                if len(body) > 2 and _NEXT_HEADING.match(nxt) and len(nxt.strip()) <= 40:
                    break
                if nxt.strip():
                    body.append(nxt.rstrip())
                j += 1
            if len(body) > 2:
                found.append({"src": name or "첨부", "lines": body})
            i = j
        if found:  # 한 파일에서 찾았으면 다른 파일(같은 내용 중복이 많음)은 안 봄
            break
    return found


def excerpt_section(text: str, kind: str) -> dict | None:
    found = excerpt_sections(text, kind)
    return found[0] if found else None


# ── 투찰 참고 ───────────────────────────────────────────────────

def _cached_fetch(text_dir: Path) -> Callable:
    """계산기 불러오기(`bidcalc.lookup`)가 공고문을 다시 받지 않게 — 폴더에 읽어 둔 글이 있으면 그걸 쓴다."""
    def fetch(session, att, timeout=30.0):
        name = attachments.original_name(attachments.AttachmentText(
            seq=att.get("seq", ""), file_name=att.get("file_name", ""), url=att.get("url", ""), ext=att.get("ext", "")))
        cached = text_dir / f"{name}.txt"
        if cached.exists():
            return attachments.AttachmentText(seq="", file_name=name, url="", ext="", text=cached.read_text(encoding="utf-8"))
        return attachments.fetch_attachment_text(session, att, timeout=timeout)
    return fetch


def bid_reference(c, config: AppConfig, text_dir: Path, history_dir: Path) -> dict:
    """입찰(적격심사·최저가) 공고만 — 기초금액·하한율·변동폭·A값 + 추천 사정률. 협상·규격가격은 가격 비중이 작아 안 함."""
    if c.category != "입찰":
        return {"skip": f"{c.category or '이 유형'} 공고 — 투찰가 계산 대상 아님(가격 위주 입찰만)"}
    if not config.api.service_key:
        return {"skip": "나라장터 인증키가 없어 기초금액을 못 불러옴"}
    try:
        info = bidcalc.lookup(bidcalc.make_client(config.api), c.notice.notice_no, fetch=_cached_fetch(text_dir))
    except ApiError as err:
        return {"skip": f"나라장터 조회 실패 — {redact(str(err), [config.api.service_key])[:120]}"}
    keep = ("basis", "basis_source", "lower_rate", "range_lo", "range_hi", "a_value", "a_applies", "net_cost", "price_method",
            "method", "method_reason", "agency_note", "notes")
    out = {k: info.get(k) for k in keep}
    if info.get("basis") is None or info.get("lower_rate") is None:
        out["skip"] = "기초금액 또는 하한율이 아직 없음(공개 전일 수 있음) — 마감 전에 이 폴더를 다시 만들면 채워짐"
        return out
    a = (info.get("a_value") or 0) if info.get("a_applies") else 0
    out["recommend"] = bidrecommend.recommend(
        bidrecommend.history_from_dir(history_dir), work=info.get("work_type"), lo=info.get("range_lo"),
        hi=info.get("range_hi"), basis=info["basis"], lower_rate=info["lower_rate"], a=a)
    return out


# ── 요약 문서(HTML → PDF) ───────────────────────────────────────

def _esc(v: Any) -> str:
    return html.escape("" if v is None else str(v))


def _money(v: float | None) -> str:
    return "-" if v is None else f"{v:,.0f}원"


def _when(dt: datetime | None) -> str:
    return "-" if dt is None else f"{dt:%Y-%m-%d %H:%M}"


def _dday(dt: datetime, now: datetime) -> str:
    days = (dt.date() - now.date()).days
    return "마감" if days < 0 else f"D-{days}"


CSS = """
@page { size: A4; margin: 13mm 12mm; }
body { font-family: 'Malgun Gothic', 'Apple SD Gothic Neo', 'Noto Sans KR', 'WenQuanYi Zen Hei', sans-serif;
  font-size: 9.5pt; color: #1d2433; line-height: 1.5; margin: 0; }
h1 { font-size: 15pt; margin: 0 0 2px; }
h3 { font-size: 9.8pt; margin: 10px 0 4px; }
pre.excerpt { white-space: pre-wrap; font-family: inherit; font-size: 8.6pt; background: #f7f8fa; border: 1px solid #d5d9e0;
  padding: 6px 8px; margin: 3px 0; }
h2 { font-size: 10.5pt; margin: 14px 0 5px; padding-bottom: 3px; border-bottom: 1.5px solid #1d2433; }
.meta, .src { color: #6b7280; font-size: 8.3pt; }
.src { white-space: nowrap; }
table { width: 100%; border-collapse: collapse; }
td, th { border: 1px solid #d5d9e0; padding: 4px 6px; vertical-align: top; text-align: left; }
th { background: #f1f3f6; }
.kv td:first-child { width: 20%; background: #f7f8fa; font-weight: 700; }
.num { text-align: right; white-space: nowrap; }
tr.key td { font-weight: 700; }
.box { margin: 10px 0; padding: 8px 11px; border: 1px solid #d5d9e0; border-radius: 6px; background: #f7f8fa; }
.ok { color: #1f7a45; } .bad { color: #b91c1c; } .warn { color: #8a5a00; }
ul { margin: 3px 0; padding-left: 18px; }
.foot { margin-top: 14px; color: #6b7280; font-size: 8pt; border-top: 1px solid #d5d9e0; padding-top: 5px; }
"""


def _qualification_html(c) -> str:
    from .email_html import _req_blocks, qualification_parts

    q = c.qualification
    if not q.checked:
        from .report import unchecked_reason

        return f'<div class="warn">자격 미확인 — {_esc(unchecked_reason(c))}</div>'
    parts = []
    for p in qualification_parts(c):
        blocks = _req_blocks(list(p.get("missing") or []) + list(p.get("satisfied") or []))
        if blocks:
            parts.append(f'<div style="margin-top:6px;font-weight:700">{_esc(p["name"])}</div>' + "".join(blocks))
    head = (f'<div class="ok"><b>자격 충족</b></div>' if not q.missing_count
            else f'<div class="bad"><b>{_esc(q.summary)}</b> — 공동수급: {_esc(c.joint.label)}</div>')
    return head + "".join(parts)


def summary_html(c, ext: dict, bid: dict | None, now: datetime, files: list[str]) -> str:
    n = c.notice
    earliest = c.schedule.earliest
    deadline_rows = "".join(
        f'<tr class="{"key" if earliest and (label, dt) == earliest else ""}"><td>{_when(dt)}</td><td>{_esc(label)}</td>'
        f'<td class="num">{"" if label in ("공고 게시", "개찰") else _dday(dt, now)}</td></tr>'
        for label, dt in c.schedule.all_deadlines)
    submissions = "".join(
        f"<tr><td>{_esc(s['what'])}</td><td>{_esc(s['when'])}</td><td>{_esc(s['how'])}</td><td>{_esc(s['where'])}</td>"
        f'<td class="src">{_esc(s["src"])}</td></tr>' for s in ext.get("submissions", []))
    flags = "".join(
        f"<li><b>{_esc(f.get('kind'))}</b> — {_esc((f.get('text') or '')[:300])}"
        + (f" <span class=src>{_esc(f['source'])}</span>" if f.get("source") else "") + "</li>"
        for f in c.text_flags)
    tasks = "".join(f"<li>{_esc(t['text'])} <span class=src>{_esc(t['src'])}</span></li>" for t in ext.get("task_summary", []))
    docs = "".join(
        f"<tr><td>{_esc(d['name'])}</td><td>{_esc(d['form'])}</td><td>{_esc(d['qty'])}</td><td>{_esc(d['stage'])}</td>"
        f"<td>{_esc(d['prep'])}</td><td>{_esc(d['note'])}</td></tr>" for d in ext.get("documents", []))
    clauses = "".join(f"<tr><td>{_esc(k['kind'])}</td><td>{_esc(k['text'])} <span class=src>{_esc(k['src'])}</span></td></tr>"
                      for k in ext.get("clauses", []))
    region = getattr(c, "region_check", None)
    region_text = (f"{region.status}" + (f" ({'·'.join(region.required)})" if getattr(region, "required", None) else "")
                   if region and getattr(region, "status", None) else "제한 정보 없음")
    if c.region_suspect:
        region_text += f" · {c.region_suspect}"
    sync = (f"{c.sync['score'] * 100:.0f}% — {_esc(c.sync['top'][0]['title'])}" if c.sync and c.sync.get("top") and c.sync.get("score") is not None
            else "-")

    def section(title: str, body: str, empty: str) -> str:
        return f"<h2>{title}</h2>" + (body if body else f'<div class="meta">{empty}</div>')

    bid_html = _bid_html(bid)
    note = ext.get("note") or ""
    head = WORK_HEADINGS.get(n.work_type, WORK_HEADINGS["용역"])
    basis = (bid or {}).get("basis")
    basis_text = (f"{_money(basis)} ({_esc((bid or {}).get('basis_source') or '')})" if basis is not None
                  else "입찰(가격) 공고가 아니거나 아직 공개 전 — 공고문 확인")
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>{_esc(n.title)}</title><style>{CSS}</style></head><body>
<div class="meta">트랙 1 공고 요약 · 자동 생성 {now:%Y-%m-%d %H:%M} · 과업·제출서류·조항: {_esc(ext.get('source'))}{(' (' + _esc(note) + ')') if note else ''}</div>
<h1>{_esc(n.title)}</h1>
<div class="meta">{_esc(n.demand_institution or n.notice_institution)} · 공고번호 {_esc(n.notice_no)}-{_esc(n.notice_ord)}
{f' · <a href="{_esc(n.detail_url)}">나라장터 상세</a>' if n.detail_url else ''}</div>

<div class="box">
<b>마감</b> {(_esc(earliest[0]) + ' ' + _when(earliest[1]) + ' (' + _dday(earliest[1], now) + ')') if earliest else '일정 미상'}
 · <b>자격</b> {_esc(c.qualification.summary)} · <b>지역</b> {_esc(region_text)} · <b>공동수급</b> {_esc(c.joint.label)}
{(' · <b>확인 필요</b> ' + _esc(', '.join(sorted({f.get('kind', '') for f in c.text_flags})))) if c.text_flags else ''}
</div>

<h2>1. 나라장터 공고문</h2>
<div>공고 본문: {f'<a href="{_esc(n.detail_url)}">나라장터 상세</a> · ' if n.detail_url else ''}00_원문/나라장터_공고정보.txt</div>
<div>첨부 원문 (00_원문 폴더): {_esc(', '.join(files)) if files else '받은 첨부 없음 — 나라장터 상세에서 확인'}</div>

<h2>2. 공고 요약 및 검토 자료</h2>
<h3>① {_esc(head['name'])}명 및 주요 {_esc(head['content'])}</h3>
<table class="kv">
<tr><td>{_esc(head['name'])}명</td><td>{_esc(n.title)}</td></tr>
<tr><td>업무 · 유형</td><td>{_esc(n.work_type)} · {_esc(c.category)} · {_esc(n.award_method or '-')}</td></tr>
<tr><td>싱크로율</td><td>{sync}</td></tr>
</table>
{f'<ul>{tasks}</ul>' if tasks else f'<div class="meta">자동 요약 없음 — 과업지시서·제안요청서·시방서 확인</div>'}

<h3>② {_esc(head['name'])} 금액 및 공고 일정</h3>
<table class="kv">
<tr><td>추정가격</td><td>{_money(n.estimated_price)}{' (부가세 제외)' if n.estimated_price else ''}</td></tr>
<tr><td>배정예산</td><td>{_money(n.assigned_budget)}</td></tr>
<tr><td>기초금액</td><td>{basis_text}</td></tr>
<tr><td>트랙 기준 금액</td><td>{_money(n.budget)} ({_esc(n.budget_kind or '금액 모름')}) — 1억 미만</td></tr>
</table>
<table style="margin-top:6px"><tr><th style="width:24%">일시</th><th>일정 (굵은 줄 = 카드 마감 기준)</th><th style="width:9%">D-day</th></tr>{deadline_rows or '<tr><td colspan=3>일정 미상 — 공고문 확인</td></tr>'}</table>
{('<div class="meta" style="margin-top:6px"><b>제출 방법</b></div><table><tr><th>제출</th><th>일시</th><th>방법</th><th>제출처</th><th>근거</th></tr>' + submissions + '</table>') if submissions else ''}

<h3>③ 입찰 참가 자격 및 필수 조건</h3>
{_qualification_html(c)}
<table class="kv" style="margin-top:6px">
<tr><td>지역 제한</td><td>{_esc(region_text)}</td></tr>
<tr><td>공동수급</td><td>{_esc(c.joint.label)}{(' · 협정 마감 ' + _when(c.schedule.joint_agreement_deadline)) if c.schedule.joint_agreement_deadline else ''}</td></tr>
</table>
{('<div style="margin-top:8px"><b>확인 필요</b> (판정 안 함 — 원문 확인)</div><ul>' + flags + '</ul>') if flags else ''}

<h3>④ 제출 서류 목록</h3>
{('<table><tr><th style="width:36%">서류</th><th style="width:13%">서식</th><th style="width:8%">수량</th><th style="width:13%">제출 단계</th><th style="width:10%">준비</th><th>주의</th></tr>' + docs + '</table><div class=meta>체크리스트 파일(02)에 같은 목록이 있습니다.</div>') if docs else _excerpt_html(ext, 'documents', '제출서류 절을 찾지 못함 — 공고문 확인')}

<h3>⑤ 기타 유의사항 및 제한 조건</h3>
{f'<table style="margin-top:6px"><tr><th style="width:12%">구분</th><th>내용</th></tr>{clauses}</table>' if clauses else _excerpt_html(ext, 'notes', '유의사항 절을 찾지 못함 — 공고문 확인')}

<h2>3. 투찰 참고</h2>
{bid_html}

<div class="foot">자동 생성 자료입니다. 판정·일정은 근거 위치의 원문과 대조하세요. 투찰 금액 결정·전자서명·서류 날인은 담당자가 직접 합니다.</div>
</body></html>"""


WORK_HEADINGS = {"공사": {"name": "공사", "content": "공사 내용"}, "용역": {"name": "용역", "content": "과업 내용"},
                 "물품": {"name": "물품", "content": "규격·납품 내용"}}


def _excerpt_html(ext: dict, kind: str, empty: str) -> str:
    found = (ext.get("excerpts") or {}).get(kind) or []
    if not found:
        return f'<div class="meta">{_esc(empty)}</div>'
    return (f'<div class="meta">원문 발췌 ({_esc(found[0]["src"])}) — 자동 정리 없이 원문 그대로</div>'
            + "".join(f'<pre class="excerpt">{_esc(chr(10).join(f["lines"]))}</pre>' for f in found))


def _bid_html(bid: dict | None) -> str:
    if not bid:
        return '<div class="meta">-</div>'
    lines = []
    if bid.get("basis") is not None:
        rng = (f"{bid['range_lo']:+g}% ~ {bid['range_hi']:+g}%" if bid.get("range_lo") is not None and bid.get("range_hi") is not None
               else "-")
        a = f"{bid['a_value']:,.0f}원{'' if bid.get('a_applies') else ' (적용 표시 없음)'}" if bid.get("a_value") else "-"
        lines.append(f"<table class=kv><tr><td>기초금액</td><td>{_money(bid['basis'])} ({_esc(bid.get('basis_source'))})</td></tr>"
                     f"<tr><td>낙찰하한율 · 변동폭</td><td>{_esc(bid.get('lower_rate'))}% · {rng}</td></tr>"
                     f"<tr><td>A값</td><td>{a}</td></tr>"
                     f"<tr><td>예비가격 방식</td><td>{_esc(bid.get('agency_note') or bid.get('method_reason') or '-')}</td></tr></table>")
    if bid.get("skip"):
        lines.append(f'<div class="warn" style="margin-top:4px">{_esc(bid["skip"])}</div>')
    rec = bid.get("recommend")
    if rec and rec.get("options"):
        rows = "".join(
            f"<tr><td>{_esc(o['name'])}</td><td class=num>{o['rate']:.3f}%</td><td class=num>{o['wins']}건 ({o['win_ratio']:.0%})</td>"
            f"<td class=num>{o['valid']}건 ({o['valid_ratio']:.0%})</td><td class=num>{_money(o.get('bid'))}</td></tr>"
            for o in rec["options"])
        cond = " · ".join(rec.get("conditions") or []) or "조건 없음"
        lines.append(f'<div style="margin-top:6px"><b>추천 사정률</b> — 비슷한 과거 공고 {rec["samples"]}건({_esc(cond)})</div>'
                     f"<table><tr><th>구분</th><th>가정 사정률</th><th>과거 1순위였을 공고</th><th>유효</th>"
                     f"<th>투찰금액(끝자리 처리 전)</th></tr>{rows}</table>")
    if rec:
        lines += [f'<div class="warn">※ {_esc(x)}</div>' for x in rec.get("notes", [])]
    lines += [f'<div class="meta">{_esc(x)}</div>' for x in (bid.get("notes") or [])]
    return "".join(lines) or '<div class="meta">-</div>'


def find_browser() -> str | None:
    """PDF로 바꿀 브라우저 — 환경변수 TRACK1_BROWSER, 엣지·크롬(윈도우 기본 위치 포함), 리눅스 크로미움 순."""
    env = os.environ.get("TRACK1_BROWSER", "").strip()
    if env:
        return env
    for name in ("msedge", "chrome", "google-chrome", "chromium", "chromium-browser"):
        found = shutil.which(name)
        if found:
            return found
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    ]
    pw = Path("/opt/pw-browsers")  # 리눅스(개발·Actions) — 화면 없는 인쇄 전용 빌드가 가장 잘 됨
    candidates += sorted(str(p) for p in pw.glob("chromium_headless_shell-*/chrome-headless-shell-*/chrome-headless-shell"))
    candidates += sorted(str(p) for p in pw.glob("chromium-*/chrome-linux*/chrome"))
    return next((p for p in candidates if Path(p).exists()), None)


def html_to_pdf(html_path: Path, pdf_path: Path, browser: str | None) -> str | None:
    """브라우저의 '인쇄 → PDF'(headless)로 PDF를 만든다. 실패하면 이유를 돌려준다(HTML은 남음)."""
    if not browser:
        return "PDF로 바꿀 브라우저(엣지·크롬)를 찾지 못함 — HTML 파일을 열어 인쇄하세요"
    with tempfile.TemporaryDirectory() as profile:
        # --disable-dev-shm-usage: 공유 메모리가 작은 환경(컨테이너)에서 인쇄가 멈추는 것 방지
        cmd = [browser, "--headless", "--disable-gpu", "--no-sandbox", "--disable-dev-shm-usage", f"--user-data-dir={profile}",
               "--no-pdf-header-footer", f"--print-to-pdf={pdf_path}", html_path.resolve().as_uri()]
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=90)
        except (OSError, subprocess.SubprocessError) as err:
            return f"PDF 변환 실패({err.__class__.__name__}) — HTML 파일을 열어 인쇄하세요"
    return None if pdf_path.exists() and pdf_path.stat().st_size > 0 else "PDF 변환 결과가 비어 있음"


# ── 체크리스트·일정·관리대장 ────────────────────────────────────

def _header(ws, columns: list[str], widths: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for i, (name, width) in enumerate(zip(columns, widths), 1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D2433")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def write_checklist(path: Path, c, ext: dict) -> bool:
    """제출서류 체크리스트. 담당자가 상태를 적는 파일이라 이미 있으면 그대로 둔다(False)."""
    if path.exists():
        return False
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "제출서류"
    columns = ["서류", "서식", "수량", "제출 단계", "준비", "주의", "근거", "담당", "상태", "메모"]
    _header(ws, columns, [46, 14, 9, 14, 11, 28, 24, 9, 10, 24])
    docs = ext.get("documents") or []
    status = DataValidation(type="list", formula1='"미착수,준비중,완료,해당없음"', allow_blank=True)
    ws.add_data_validation(status)
    for r, d in enumerate(docs, 2):
        for col, value in enumerate([d["name"], d["form"], d["qty"], d["stage"], d["prep"], d["note"], d["src"], "",
                                     "해당없음" if d["prep"] == "해당 시" else "미착수", ""], 1):
            ws.cell(row=r, column=col, value=value).alignment = Alignment(wrap_text=True, vertical="top")
        status.add(f"I{r}")
    if not docs:
        ws.cell(row=2, column=1, value=f"자동 추출 없음({ext.get('note') or '규칙 판정만'}) — 공고문 '제출서류' 절을 보고 채우세요")
    ws2 = wb.create_sheet("제출 방법")
    _header(ws2, ["제출", "일시", "방법", "제출처", "근거"], [34, 30, 12, 34, 24])
    for r, s in enumerate(ext.get("submissions") or [], 2):
        for col, value in enumerate([s["what"], s["when"], s["how"], s["where"], s["src"]], 1):
            ws2.cell(row=r, column=col, value=value).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return True


def write_ics(path: Path, c, now: datetime) -> int:
    """마감 일정 → 캘린더 파일(하루 전 알림). 공고 게시는 뺌. 넣은 일정 수."""
    def stamp(dt: datetime) -> str:
        return dt.strftime("%Y%m%dT%H%M00")

    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//jiil//track1//KO", "CALSCALE:GREGORIAN"]
    count = 0
    for i, (label, dt) in enumerate(c.schedule.all_deadlines):
        if label == "공고 게시":
            continue
        uid = f"{c.notice.notice_no}-{c.notice.notice_ord}-{i}@jiil"
        lines += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{now.strftime('%Y%m%dT%H%M%S')}",
                  f"DTSTART;TZID=Asia/Seoul:{stamp(dt)}", f"DTEND;TZID=Asia/Seoul:{stamp(dt + timedelta(minutes=30))}",
                  f"SUMMARY:[입찰] {label} — {c.notice.title[:40]}", f"DESCRIPTION:공고번호 {c.notice.notice_no}"]
        if label != "개찰":
            lines += ["BEGIN:VALARM", "TRIGGER:-P1D", "ACTION:DISPLAY", f"DESCRIPTION:내일 {label}", "END:VALARM"]
        lines.append("END:VEVENT")
        count += 1
    lines.append("END:VCALENDAR")
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")
    return count


def ledger_row(c, folder: Path, root: Path, now: datetime) -> dict:
    earliest = c.schedule.earliest
    return {
        "공고번호": c.notice.notice_no, "차수": c.notice.notice_ord, "등록일": f"{now:%Y-%m-%d}", "공고명": c.notice.title,
        "발주기관": c.notice.demand_institution or c.notice.notice_institution or "", "업무": c.notice.work_type,
        "유형": c.category or "", "금액": c.notice.budget, "금액 종류": c.notice.budget_kind or "",
        "마감": f"{earliest[1]:%Y-%m-%d %H:%M} {earliest[0]}" if earliest else "일정 미상",
        "자격": c.qualification.summary, "지역": getattr(c.region_check, "status", "") or "",
        "공동수급": c.joint.label, "확인 필요": ", ".join(sorted({f.get("kind", "") for f in c.text_flags})),
        "폴더": str(folder.relative_to(root)) if folder.is_relative_to(root) else str(folder),
    }


def update_ledger(path: Path, rows: list[dict]) -> None:
    """공고번호로 찾아 시스템 칸만 고치고(차수가 바뀌면 정정공고), 없으면 아래에 더함. 담당자 칸은 건드리지 않는다.
    등록일은 처음 넣은 날 그대로."""
    from openpyxl import Workbook, load_workbook

    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "공고"
        _header(ws, SYSTEM_COLUMNS + STAFF_COLUMNS,
                [16, 6, 11, 44, 22, 7, 10, 14, 9, 22, 26, 8, 14, 18, 44, 10, 8, 10, 24])
    header = [cell.value for cell in ws[1]]
    col = {name: header.index(name) + 1 for name in header if name}
    for name in SYSTEM_COLUMNS + STAFF_COLUMNS:  # 예전 파일에 없는 칸은 오른쪽에 더함
        if name not in col:
            col[name] = ws.max_column + 1
            ws.cell(row=1, column=col[name], value=name)
    existing = {ws.cell(row=r, column=col["공고번호"]).value: r for r in range(2, ws.max_row + 1)}
    for row in rows:
        r = existing.get(row["공고번호"])
        if r is None:
            r = ws.max_row + 1
            existing[row["공고번호"]] = r
        for name in SYSTEM_COLUMNS:
            if name == "등록일" and ws.cell(row=r, column=col[name]).value:
                continue
            cell = ws.cell(row=r, column=col[name], value=row[name])
            if name == "금액":
                cell.number_format = "#,##0"
    wb.save(path)


# ── 묶어서 ─────────────────────────────────────────────────────

@dataclass
class Track1Options:
    root: Path
    max_budget: float = DEFAULT_MAX_BUDGET
    llm: LlmSettings | None = None
    browser: str | None = None
    history_dir: Path = Path("output")
    pdf: bool = True


@dataclass
class BuildResult:
    folder: Path
    pdf_error: str | None = None
    checklist_written: bool = False
    extraction: str = ""
    errors: list[str] = field(default_factory=list)


NOTICE_LABELS = {
    "title": "공고명", "notice_no": "공고번호", "notice_ord": "차수", "work_type": "업무", "notice_institution": "공고기관",
    "demand_institution": "수요기관", "award_method": "낙찰방법", "contract_method": "계약방법", "bid_method": "입찰방식",
    "estimated_price": "추정가격", "assigned_budget": "배정예산", "posted_at": "공고 게시", "qualification_deadline": "입찰참가자격등록 마감",
    "bid_deadline": "입찰서 제출 마감", "opening_at": "개찰", "joint_method_name": "공동수급", "joint_agreement_deadline": "공동수급협정 마감",
    "estimate_price_method": "예가방법", "industry_text": "업종 제한", "product_class_name": "세부품명", "product_class_no": "세부품명번호",
    "detail_url": "나라장터 상세 주소",
}


def write_notice_info(path: Path, n) -> None:
    """나라장터 공고 본문 값(첨부가 아닌 공고 상세) — 사람이 읽는 칸 + 원본 응답 전체. 개인 연락처는 수집 때 이미 뺌."""
    lines = ["나라장터 공고 정보 (입찰공고정보서비스 응답)", ""]
    for key, label in NOTICE_LABELS.items():
        value = getattr(n, key, None)
        if value not in (None, "", []):
            lines.append(f"{label}: {f'{value:,.0f}원' if isinstance(value, float) else value}")
    lines += ["", "첨부 목록:"] + [f"  - {a.get('file_name') or '(이름 없음)'}  {a.get('url', '')}" for a in n.attachments]
    lines += ["", "── 원본 응답 ──", json.dumps(n.raw, ensure_ascii=False, indent=1)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_folder(c, config: AppConfig, options: Track1Options, now: datetime, llm_client: Any = None) -> BuildResult:
    folder = folder_for(options.root, c, now)
    raw = folder / "00_원문"
    raw.mkdir(parents=True, exist_ok=True)
    write_notice_info(raw / "나라장터_공고정보.txt", c.notice)
    result = BuildResult(folder=folder)
    if options.llm and llm_client is not None:
        ext = extract_with_llm(llm_client, options.llm, c)
        if ext.get("source") != "Claude":
            ext = {**rule_extraction(c, ext.get("note", "")), "note": ext.get("note", "")}
    else:
        ext = rule_extraction(c, "Claude API 키가 없어 과업 요약은 '사업 범위' 목록, 제출서류·유의사항은 원문 발췌")
    result.extraction = ext.get("source", "")
    bid = bid_reference(c, config, raw / "텍스트", options.history_dir)
    files = sorted(p.name for p in raw.iterdir() if p.is_file() and p.name != "나라장터_공고정보.txt")

    page = folder / "01_공고요약.html"
    page.write_text(summary_html(c, ext, bid, now, files), encoding="utf-8")
    if options.pdf:
        result.pdf_error = html_to_pdf(page, folder / "01_공고요약.pdf", options.browser)
    try:
        result.checklist_written = write_checklist(folder / "02_제출서류_체크리스트.xlsx", c, ext)
    except ImportError:
        result.errors.append("openpyxl이 없어 체크리스트를 못 만듦 — pip install -r requirements.txt")
    write_ics(folder / "일정.ics", c, now)
    record = {
        "공고": {"번호": c.notice.notice_no, "차수": c.notice.notice_ord, "공고명": c.notice.title, "금액": c.notice.budget,
                 "금액 종류": c.notice.budget_kind, "유형": c.category, "업무": c.notice.work_type},
        "일정": [{"이름": label, "일시": f"{dt:%Y-%m-%d %H:%M}"} for label, dt in c.schedule.all_deadlines],
        "자격": c.qualification.summary, "공동수급": c.joint.label, "확인 필요": c.text_flags,
        "자동 추출": ext, "투찰 참고": bid, "만든 시각": f"{now:%Y-%m-%d %H:%M}",
    }
    (folder / "추출결과.json").write_text(json.dumps(record, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return result


def build_folders(candidates: list, config: AppConfig, options: Track1Options, now: datetime | None = None,
                  progress: Callable[[int, int], None] | None = None) -> list[BuildResult]:
    """트랙 1 공고마다 폴더를 만들고 관리대장을 고친다. 첨부 원본은 `originals_hook`로 수집 때 이미 받아 둔 것을 쓴다."""
    now = now or datetime.now()
    track = [c for c in candidates if is_track1(c, options.max_budget)]
    llm_client = None
    if options.llm:
        from .llm_similarity import make_client

        llm_client = make_client()
    results, rows = [], []
    for i, c in enumerate(track):
        if progress:
            progress(i, len(track))
        try:
            res = build_folder(c, config, options, now, llm_client)
        except OSError as err:  # NAS가 끊겨도 다른 공고는 계속
            log.warning("트랙 1 폴더 만들기 실패 [%s]: %s", c.notice.notice_no, err)
            continue
        results.append(res)
        rows.append(ledger_row(c, res.folder, options.root, now))
        log.info("트랙 1 폴더: %s (요약 %s%s)", res.folder, res.extraction,
                 f", PDF 안 됨: {res.pdf_error}" if res.pdf_error else "")
    if rows:
        try:
            update_ledger(options.root / "입찰" / LEDGER_NAME, rows)
        except ImportError:
            log.warning("openpyxl이 없어 관리대장을 못 고침 — pip install -r requirements.txt")
        except PermissionError:
            log.warning("관리대장이 열려 있어 저장 못 함(엑셀을 닫고 다시 실행): %s", options.root / "입찰" / LEDGER_NAME)
    return results


def originals_hook(options: Track1Options, now: datetime) -> Callable:
    """`attachments.save_attachment_texts(originals=…)` — 트랙 1 공고만 첨부 원본을 폴더 00_원문에 남긴다."""
    def hook(c) -> Path | None:
        return folder_for(options.root, c, now) / "00_원문" if is_track1(c, options.max_budget) else None
    return hook


def drop_unbuilt(kept: list, before: list, options: Track1Options, now: datetime) -> None:
    """첨부를 읽은 뒤 마감 지남·지역 미달로 빠진 트랙 1 공고 — 원본만 받아 둔 폴더를 지운다(요약을 만든 적 있는 폴더는 둠)."""
    kept_ids = {id(c) for c in kept}
    for c in before:
        if id(c) in kept_ids or not is_track1(c, options.max_budget):
            continue
        folder = folder_for(options.root, c, now)
        if folder.exists() and not (folder / "01_공고요약.html").exists():
            shutil.rmtree(folder, ignore_errors=True)


def options_from_env(root: str | Path | None = None, llm: bool = True, pdf: bool = True) -> Track1Options | None:
    """환경변수로 옵션 — TRACK1_ROOT(필수), TRACK1_MAX_BUDGET, TRACK1_BROWSER, BID_HISTORY_DIR. 루트가 없으면 None(끔)."""
    root = root or os.environ.get("TRACK1_ROOT", "").strip()
    if not root:
        return None
    return Track1Options(
        root=Path(root), max_budget=float(os.environ.get("TRACK1_MAX_BUDGET") or DEFAULT_MAX_BUDGET),
        llm=LlmSettings.from_env() if llm else None, browser=find_browser() if pdf else None,
        history_dir=Path(os.environ.get("BID_HISTORY_DIR") or os.environ.get("OUTPUT_DIR") or "output"), pdf=pdf)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="트랙 1(1억 미만) 공고를 NAS 폴더로 — 수집·자격판정 후 공고마다 자료 생성")
    parser.add_argument("--root", help="폴더를 만들 곳(NAS 공유 폴더). 없으면 환경변수 TRACK1_ROOT")
    parser.add_argument("--days", type=int, default=1, help="최근 며칠 게시분(기본 1)")
    parser.add_argument("--categories", default="", help="유형 쉼표 목록(협상,규격가격동시입찰,입찰) — 비우면 전부")
    parser.add_argument("--no-llm", action="store_true", help="Claude API를 쓰지 않음(규칙 판정만)")
    parser.add_argument("--no-pdf", action="store_true", help="PDF로 바꾸지 않음(HTML만)")
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)-7s %(message)s", datefmt="%H:%M:%S")

    options = options_from_env(args.root, llm=not args.no_llm, pdf=not args.no_pdf)
    if options is None:
        print("폴더를 만들 곳이 없습니다 — --root 또는 환경변수 TRACK1_ROOT", file=sys.stderr)
        return 2
    try:
        config = load_config()
    except ConfigError as err:
        print(f"설정 오류: {err}", file=sys.stderr)
        return 2
    if not config.api.service_key:
        print("NARA_SERVICE_KEY 환경변수가 필요합니다.", file=sys.stderr)
        return 1
    from . import scope
    from .attachments import save_attachment_texts
    from .pipeline import drop_expired, drop_region_unmet, run
    from .sync import PastIndex

    now = datetime.now()
    config.lookback_days = args.days
    try:
        candidates, stats, _ = run(config, now, scope.parse_categories(args.categories))
    except (ApiError, ValueError) as err:
        print(redact(str(err), [config.api.service_key]), file=sys.stderr)
        return 1
    track = [c for c in candidates if is_track1(c, options.max_budget)]
    print(f"후보 {len(candidates)}건 중 트랙 1(금액 {options.max_budget:,.0f}원 미만 본공고) {len(track)}건")
    save_attachment_texts(
        track, config.output_dir, held_codes=qualify.load_held_codes(config.held_raw),
        held_code_names=qualify.load_held_code_names(config.held_raw), now=now, code_names=config.code_names,
        held_raw=config.held_raw, client=DataGoKrClient(config.api), originals=originals_hook(options, now))
    before = list(track)
    track = drop_region_unmet(drop_expired(track, now, stats), stats)
    drop_unbuilt(track, before, options, now)
    past = PastIndex.load()
    if past.available:
        for c in track:
            c.sync = past.score(c.notice.title, c.task_text or c.attachment_text)
    if not options.llm and not args.no_llm:
        print("ANTHROPIC_API_KEY가 없어 과업 요약·제출서류는 규칙으로만 만듭니다")
    results = build_folders(track, config, options, now)
    for res in results:
        print(f"  {res.folder}" + (f"  (PDF 안 됨: {res.pdf_error})" if res.pdf_error else ""))
    print(f"폴더 {len(results)}개 · 관리대장 {options.root / '입찰' / LEDGER_NAME}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
