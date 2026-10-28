"""트랙 1 공고 폴더(nego/track1.py, 2026-10-28) — 나라장터·Claude·브라우저는 부르지 않고 폴더 내용만 확인한다."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from nego import track1  # noqa: E402
from nego.attachments import AttachmentText, save_original  # noqa: E402
from nego.config import load_config  # noqa: E402
from nego.models import notice_from_raw  # noqa: E402
from nego.pipeline import RunStats, build_candidates  # noqa: E402
from tests import fixtures  # noqa: E402

NOW = datetime(2026, 9, 2, 9, 0)


def _candidates():
    config = load_config()
    config.api.service_key = ""
    config.screen.keywords = ["전시관"]
    raws = [
        fixtures.notice("R26BK00000001", title="○○전시관 설계: 용역/1차", presmptPrce="24660000", asignBdgtAmt="27126000"),
        fixtures.notice("R26BK00000002", title="○○전시관 제작설치"),  # 5억 — 트랙 1 아님
    ]
    cands = build_candidates([notice_from_raw(r, "용역") for r in raws], config, {}, {}, datetime(2026, 9, 1),
                             RunStats())
    return config, {c.notice.notice_no: c for c in cands}


class FakeLlm:
    """Claude 응답 흉내 — beta.messages.create가 JSON 텍스트 블록 하나를 돌려준다."""

    def __init__(self, data):
        self.data = data
        self.calls = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))

    def _create(self, **kw):
        self.calls.append(kw)
        return SimpleNamespace(stop_reason="end_turn", model="claude-test",
                               content=[SimpleNamespace(type="text", text=json.dumps(self.data, ensure_ascii=False))])


EXTRACTED = {
    "task_summary": [{"text": "말산업 체험공간 207㎡ 전시 설계", "src": "제안요청서 · Ⅰ-2"}],
    "submissions": [{"what": "입찰참가신청서류", "when": "2026. 9. 20. 09:00~11:00", "how": "이메일", "where": "행정지원팀",
                     "src": "공고문 · 3-다"}],
    "documents": [
        {"name": "입찰참가신청서", "form": "별지1", "qty": "1부", "stage": "참가신청", "prep": "서식 작성", "note": "", "src": "공고문 6"},
        {"name": "공동수급협정서", "form": "별지18", "qty": "1부", "stage": "참가신청", "prep": "해당 시", "note": "공동수급 시",
         "src": "공고문 6"},
    ],
    "clauses": [{"kind": "가격", "text": "예정가격 80% 미만은 80%로 환산", "src": "제안요청서 Ⅴ-6"}],
    "proposal_outline": [],
}


class Track1Test(unittest.TestCase):
    def setUp(self):
        self.config, self.cands = _candidates()
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.options = track1.Track1Options(root=self.root, pdf=False, history_dir=self.root / "없음")

    def tearDown(self):
        self.tmp.cleanup()

    def test_track_and_folder_name(self):
        small, big = self.cands["R26BK00000001"], self.cands["R26BK00000002"]
        self.assertTrue(track1.is_track1(small))
        self.assertFalse(track1.is_track1(big))
        folder = track1.folder_for(self.root, small, NOW)
        self.assertEqual(folder.parent, self.root / "입찰" / "2026")
        self.assertEqual(folder.name, "0901_R26BK00000001_○○전시관 설계_ 용역_1차")  # 게시일 9/1, 파일 이름에 못 쓰는 글자는 _

    def test_build_without_llm(self):
        results = track1.build_folders(list(self.cands.values()), self.config, self.options, NOW)
        self.assertEqual(len(results), 1)
        folder = results[0].folder
        for name in ("01_공고요약.html", "02_제출서류_체크리스트.xlsx", "일정.ics", "추출결과.json"):
            self.assertTrue((folder / name).exists(), name)
        page = (folder / "01_공고요약.html").read_text(encoding="utf-8")
        self.assertIn("24,660,000원 (추정가격)", page)
        self.assertIn("협상 공고 — 투찰가 계산 대상 아님", page)
        self.assertIn("Claude API 키가 없어", page)
        self.assertNotIn("배점", page)
        # 10/28 필수 자료: 공고문 + 요약 다섯 항목
        for heading in ("1. 나라장터 공고문", "용역명 및 주요 과업 내용", "용역 금액 및 공고 일정", "입찰 참가 자격 및 필수 조건",
                        "제출 서류 목록", "기타 유의사항 및 제한 조건"):
            self.assertIn(heading, page)
        self.assertIn("27,126,000원", page)  # 추정가격·배정예산 둘 다
        info = (folder / "00_원문" / "나라장터_공고정보.txt").read_text(encoding="utf-8")
        self.assertIn("공고명: ○○전시관 설계: 용역/1차", info)
        self.assertIn("배정예산: 27,126,000원", info)
        self.assertNotIn("홍길동", info)  # 개인 연락처는 수집 때 빠짐
        ics = (folder / "일정.ics").read_text(encoding="utf-8")
        self.assertIn("SUMMARY:[입찰] 자격등록 마감", ics)
        self.assertNotIn("공고 게시", ics)
        self.assertIn("TRIGGER:-P1D", ics)

    def test_llm_extraction_fills_checklist(self):
        from openpyxl import load_workbook

        llm = FakeLlm(EXTRACTED)
        self.options.llm = track1.LlmSettings(model="claude-test")
        c = self.cands["R26BK00000001"]
        c.attachment_text = "=== 공고문.hwp ===\n입찰참가신청서 …"
        res = track1.build_folder(c, self.config, self.options, NOW, llm)
        self.assertEqual(res.extraction, "Claude")
        sent = llm.calls[0]
        self.assertEqual(sent["output_config"]["format"]["type"], "json_schema")
        self.assertIn("입찰참가신청서 …", sent["messages"][0]["content"])
        ws = load_workbook(res.folder / "02_제출서류_체크리스트.xlsx")["제출서류"]
        rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]
        self.assertEqual(rows[0][:3], ["입찰참가신청서", "별지1", "1부"])
        self.assertEqual(rows[0][8], "미착수")
        self.assertEqual(rows[1][8], "해당없음")  # 조건부 서류
        page = (res.folder / "01_공고요약.html").read_text(encoding="utf-8")
        for text in ("말산업 체험공간 207㎡", "이메일", "80%로 환산"):
            self.assertIn(text, page)

    def test_excerpt_sections_without_llm(self):
        text = ("=== 제안요청서.hwp ===\n6. 제출서류\n가. 제안서 10부\n=== 입찰공고문.pdf ===\n1. 입찰에 부치는 사항\n사업명: 가\n"
                "6. 제출서류 : 서식은 제안요청서 참조\n① 입찰참가신청서(별지1) 1부\n② 사업자등록증 사본\n③ 법인인감증명서\n"
                "7. 협상적격자 선정\n가. 기술 90\n14. 기타사항\n가. 하도급 불허\n나. 공채 매입\n다. 인지세\n")
        docs = track1.excerpt_section(text, "documents")
        self.assertEqual(docs["src"], "입찰공고문.pdf")  # 공고문 파일을 먼저
        self.assertEqual(docs["lines"][-1], "③ 법인인감증명서")  # 다음 큰 제목(7.) 전까지
        self.assertIn("가. 하도급 불허", track1.excerpt_section(text, "notes")["lines"])
        text += "15. 입찰 무효\n가. 서류 위조\n나. 기한 경과\n"
        notes = track1.excerpt_sections(text, "notes", 3)
        self.assertEqual([n["lines"][0] for n in notes], ["14. 기타사항", "15. 입찰 무효"])
        self.assertIsNone(track1.excerpt_section("본문뿐", "documents"))

    def test_rerun_keeps_staff_edits(self):
        from openpyxl import load_workbook

        c = self.cands["R26BK00000001"]
        track1.build_folders([c], self.config, self.options, NOW)
        folder = track1.folder_for(self.root, c, NOW)
        checklist = folder / "02_제출서류_체크리스트.xlsx"
        wb = load_workbook(checklist)
        wb.active["J2"] = "담당자 메모"
        wb.save(checklist)
        ledger = self.root / "입찰" / track1.LEDGER_NAME
        wb = load_workbook(ledger)
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        ws.cell(row=2, column=header.index("참여 여부") + 1, value="참가")
        wb.save(ledger)

        track1.build_folders([c], self.config, self.options, datetime(2026, 9, 3, 9, 0))
        self.assertEqual(load_workbook(checklist).active["J2"].value, "담당자 메모")
        ws = load_workbook(ledger).active
        header = [cell.value for cell in ws[1]]
        self.assertEqual(ws.max_row, 2, "같은 공고는 한 줄")
        self.assertEqual(ws.cell(row=2, column=header.index("참여 여부") + 1).value, "참가")
        self.assertEqual(ws.cell(row=2, column=header.index("등록일") + 1).value, "2026-09-02", "등록일은 처음 날짜 그대로")
        self.assertEqual(ws.cell(row=2, column=header.index("금액") + 1).value, 24_660_000)

    def test_save_original_and_cleanup(self):
        raw = self.root / "00_원문"
        result = AttachmentText(seq="1", file_name="공고문.pdf", url="u", ext="pdf", text="본문", data=b"%PDF-1.4")
        path = save_original(result, raw)
        self.assertEqual(path.read_bytes(), b"%PDF-1.4")
        self.assertEqual((raw / "텍스트" / "공고문.pdf.txt").read_text(encoding="utf-8"), "본문")
        nameless = AttachmentText(seq="2", file_name="", url="u", ext="", data=b"%PDF-1.4", skipped=True, error="x")
        self.assertEqual(save_original(nameless, raw).name, "첨부2.pdf")

        c = self.cands["R26BK00000001"]
        hook = track1.originals_hook(self.options, NOW)
        self.assertIsNone(hook(self.cands["R26BK00000002"]))
        hook(c).mkdir(parents=True)
        track1.drop_unbuilt([], [c], self.options, NOW)  # 지역 미달 등으로 빠지면 원본만 받은 폴더는 지움
        self.assertFalse(track1.folder_for(self.root, c, NOW).exists())

    def test_bid_reference_only_for_price_bids(self):
        c = self.cands["R26BK00000001"]
        self.assertIn("투찰가 계산 대상 아님", track1.bid_reference(c, self.config, self.root, self.root)["skip"])
        c.category = "입찰"
        self.assertIn("인증키", track1.bid_reference(c, self.config, self.root, self.root)["skip"])


if __name__ == "__main__":
    unittest.main()
